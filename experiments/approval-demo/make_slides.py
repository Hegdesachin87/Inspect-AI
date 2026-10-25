#!/usr/bin/env python3
"""Generate approval-demo.pptx — a deck explaining the approval experiment.

Usage: python make_slides.py [RESULTS_JSON]   (run with .venv-pptx)
"""

import json
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

INK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x8A, 0xBE, 0xB7)
MUTED = RGBColor(0x5A, 0x5A, 0x72)


def add_slide(prs, title, bullets=None, code=None, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(0.9))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.runs[0].font.size = Pt(34)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = INK

    if bullets:
        b = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(8.8), Inches(4.4))
        tf = b.text_frame
        tf.word_wrap = True
        for i, item in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = item
            para.space_after = Pt(10)
            for run in para.runs:
                run.font.size = Pt(19)
                run.font.color.rgb = INK
    if code:
        c = slide.shapes.add_textbox(Inches(0.6), Inches(4.9), Inches(8.8), Inches(2.1))
        tf = c.text_frame
        tf.word_wrap = False
        for i, line in enumerate(code):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            for run in para.runs:
                run.font.name = "Menlo"
                run.font.size = Pt(13)
                run.font.color.rgb = MUTED
    if note:
        n = slide.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(8.8), Inches(0.5))
        n.text_frame.paragraphs[0].text = note
        for run in n.text_frame.paragraphs[0].runs:
            run.font.size = Pt(13)
            run.font.italic = True
            run.font.color.rgb = ACCENT
    return slide


def results_lines(path):
    with open(path) as f:
        r = json.load(f)
    lines = [f"Model: {r['model']}  ·  samples: {r['samples']}", ""]
    for k, v in r["counts"].items():
        lines.append(f"{k:<34s} {v}")
    lines.append("")
    for row in r["decisions"][:14]:
        lines.append(f"  {row['tool']:<8s} {row['command'][:44]:<46s} -> {row['decision']}")
    return lines


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_slide(
        prs,
        "The eval that stops and asks",
        bullets=[
            "Tool-call approval with Inspect AI",
            "An agent demonstrates bash and python commands; the policy decides what may run",
            "Safe commands pass automatically - everything else escalates to a human",
        ],
        note="experiments/approval-demo · docs/small-inspect-experiments.md, experiment 2",
    )
    add_slide(
        prs,
        "The question",
        bullets=[
            "What actually happens when an agent tries a destructive command mid-evaluation?",
            "Can an allowlist handle the safe cases without pestering the human?",
            "What does the agent do after a rejection?",
        ],
    )
    add_slide(
        prs,
        "The approval model",
        bullets=[
            "Approvers are chained; the first one bound to a tool decides",
            "Decisions: approve · reject · escalate (next approver) · terminate",
            "The example ships two custom approvers: bash_allowlist and python_allowlist (~40 lines each)",
        ],
        note="approval.qmd: human and auto are built in; custom @approvers are plain functions",
    )
    add_slide(
        prs,
        "The policy",
        bullets=[
            "bash_allowlist: ls, echo, cat approved; rm escalates",
            "python_allowlist: print/math.factorial approved; shutil.rmtree escalates",
            "human (interactive demo): you approve or reject at the terminal, mid-run",
            "auto with decision: reject (recorded run): escalations are rejected without a TTY",
        ],
        code=["approvers:", "  - name: bash_allowlist", "      tools: '*bash*'", "      allowed_commands: [ls, echo, cat]", "  - name: human", "      tools: '*'"],
    )
    add_slide(
        prs,
        "Run it two ways",
        bullets=[
            "./run.sh - the interactive demo: the eval freezes at rm and rmtree and asks you",
            "./run.sh --record - unattended: escalations auto-rejected, decisions recorded in the log",
            "Both write a .eval log; every decision is inspectable afterwards",
        ],
        code=["export DEEPSEEK_API_KEY=...", "./run.sh            # interactive", "./run.sh --record   # unattended"],
        note="Docker required - the tools execute inside a sandbox container",
    )
    if len(sys.argv) > 1:
        slide = add_slide(prs, "Results (recorded run)")
        b = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(8.8), Inches(5.2))
        tf = b.text_frame
        for i, line in enumerate(results_lines(sys.argv[1])):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            for run in para.runs:
                run.font.name = "Menlo"
                run.font.size = Pt(14)
                run.font.color.rgb = INK
    add_slide(
        prs,
        "Approval vs review",
        bullets=[
            "Approval gates execution BEFORE the command runs",
            "Review judges tool calls AFTER execution - and can terminate the run (review.qmd)",
            "Both leave a decision trail in the log - governance you can audit, not just observe",
        ],
    )
    add_slide(
        prs,
        "Why it matters",
        bullets=[
            "This is the governance primitive for agentic AI - experienced, not described",
            "Policy as config: the same task runs hands-off or human-gated by swapping one YAML file",
            "The log records what was attempted, what was allowed, and what the agent did next",
        ],
        note="guide: docs/small-inspect-experiments.md",
    )

    prs.save("approval-demo.pptx")
    print("wrote approval-demo.pptx")


if __name__ == "__main__":
    main()
