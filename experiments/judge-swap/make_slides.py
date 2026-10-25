#!/usr/bin/env python3
"""Generate judge-swap.pptx — a deck explaining the judge-swap experiment.

Usage: python make_slides.py [LOG_FILE]
With a log file, the results slide carries the recorded numbers.
Run with the .venv-pptx environment (python-pptx).
"""

import sys

import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

INK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x8A, 0xBE, 0xB7)
MUTED = RGBColor(0x5A, 0x5A, 0x72)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)


def add_slide(prs, title, bullets=None, code=None, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(8.8), Inches(0.9))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.runs[0].font.size = Pt(34)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = INK

    y = 1.6
    if bullets:
        b = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(8.8), Inches(4.4))
        tf = b.text_frame
        tf.word_wrap = True
        for i, item in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = item
            para.space_after = Pt(10)
            for run in para.runs:
                run.font.size = Pt(19)
                run.font.color.rgb = INK
        y += 3.6
    if code:
        c = slide.shapes.add_textbox(Inches(0.6), Inches(4.6), Inches(8.8), Inches(2.4))
        tf = c.text_frame
        tf.word_wrap = False
        for i, line in enumerate(code):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            for run in para.runs:
                run.font.name = "Menlo"
                run.font.size = Pt(13)
                run.font.color.rgb = MUTED
    if note:
        n = slide.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(8.8), Inches(0.5))
        n.text_frame.paragraphs[0].text = note
        for run in n.text_frame.paragraphs[0].runs:
            run.font.size = Pt(13)
            run.font.italic = True
            run.font.color.rgb = ACCENT
    return slide


def results_lines(path):
    with open(path) as f:
        r = json.load(f)
    lines = [f"Evaluated model: {r['model']}  ·  samples: {r['samples']}", ""]
    for name, acc in r["accuracy"].items():
        lines.append(f"{name:<24s} accuracy = {acc}")
    d = r.get("disagreements") or []
    lines.append("")
    if d:
        lines.append(f"judges disagree with matcher on {len(d)} samples:")
        for row in d[:8]:
            cells = "  ".join(f"{v:>4s}" for v in row["verdicts"].values())
            lines.append(f"  sample {row['id']:<6s} {cells}")
    else:
        lines.append("all judges agree on every sample")
    return lines


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_slide(
        prs,
        "Same outputs, different judges",
        bullets=[
            "A judge-swap experiment with Inspect AI",
            "One set of agent transcripts, four different judges",
            "Cost: cents. Runtime: minutes.",
        ],
        note="experiments/judge-swap · docs/small-inspect-experiments.md, experiment 1",
    )
    add_slide(
        prs,
        "The question",
        bullets=[
            "When the judge is a model instead of a test, does the verdict depend on which model judges?",
            "Does a model grade its own answers differently than another model does?",
            "Fixed tests gave identical grades for identical code (our Oct 4 finding) — do LLM judges behave the same?",
        ],
    )
    add_slide(
        prs,
        "The trick: score, don't regenerate",
        bullets=[
            "Generate once with --no-score: the agent's answers are frozen in a .eval log.",
            "Re-score the same log with inspect score — each judge sees byte-identical outputs.",
            "Judge differences can never be confused with generation differences.",
        ],
        code=[
            "inspect score LOG --scorer model_graded_fact \\",
            "   --model openai/glm-5.3-flash --model-base-url ... \\",
            "   --action append --overwrite",
        ],
        note="scoring bypasses the entire generation phase",
    )
    add_slide(
        prs,
        "The four judges",
        bullets=[
            "includes — deterministic string matcher, zero model calls",
            "self-grade — the evaluated model grades its own answers",
            "cross-grade, same family — glm-5.3-flash judges glm-5.3",
            "cross-grade, previous generation — glm-4.7 judges glm-5.3",
        ],
        note="theory of mind · 5 samples · chain-of-thought solver",
    )
    add_slide(
        prs,
        "What we ran",
        bullets=[
            "Task: theory of mind (official Inspect example dataset)",
            "Model: glm-5.3 via z.ai's OpenAI-compatible endpoint",
            "Unscored generation, then four appended scorings in one log",
            "-M responses_api=false forces chat completions (z.ai has no /responses)",
        ],
        code=[
            "export ZAI_API_KEY=...",
            "./run.sh          # arm 0 + all four scorings",
            "python summarize.py logs/<newest>.eval",
        ],
    )
    if len(sys.argv) > 1:
        slide = add_slide(prs, "Results (recorded run)")
        tf = slide.shapes[-1].text_frame if slide.shapes[-1].has_text_frame else None
        b = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(8.8), Inches(5.2))
        tf = b.text_frame
        for i, line in enumerate(results_lines(sys.argv[1])):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            for run in para.runs:
                run.font.name = "Menlo"
                run.font.size = Pt(14)
                run.font.color.rgb = INK
    add_slide(
        prs,
        "What to look for",
        bullets=[
            "Where the judges disagree — and which verdict is actually right",
            "The matcher's failures on open-ended answers (it cannot judge paraphrases)",
            "Self-grade vs cross-grade: any systematic leniency toward its own answers?",
            "In the viewer: each judge's written rationale under every sample",
        ],
        note="inspect view --log-dir logs --port 7575",
    )
    add_slide(
        prs,
        "Why it matters",
        bullets=[
            "LLM judges are measurement instruments — this calibrates them against each other",
            "Self-grading is the default in Inspect's model-graded scorers; know its bias before you trust the number",
            "Cents and minutes: judge quality can be audited before any expensive benchmark run",
        ],
    )
    add_slide(
        prs,
        "Run it yourself",
        bullets=[
            "experiments/judge-swap/run.sh — end-to-end: generation + all scorings",
            "Defaults: glm-5.3 on z.ai; override MODEL/BASE_URL/KEY_VAR for DeepSeek",
            "summarize.py prints the judge matrix; make_slides.py rebuilds this deck from a log",
        ],
        code=["export ZAI_API_KEY=... && ./run.sh"],
        note="guide: docs/small-inspect-experiments.md",
    )

    prs.save("judge-swap.pptx")
    print("wrote judge-swap.pptx")


if __name__ == "__main__":
    main()
